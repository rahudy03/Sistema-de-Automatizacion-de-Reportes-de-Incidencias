import json
import os
import re
import requests
import streamlit as st
from datetime import datetime
from io import BytesIO
from openai import OpenAI

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# =========================================================
# API KEYS (desde secrets.toml)
# =========================================================
WINDY_API_KEY = st.secrets["WINDY_API_KEY"]
DEEPSEEK_API_KEY = st.secrets["DEEPSEEK_API_KEY"]
OPENROUTER_API_KEY = st.secrets["OPENROUTER_API_KEY"]


# =========================================================
# HELPERS DE CONVERSIÓN SEGURA
# =========================================================
def int_seguro(valor, default=0):
    try:
        return int(valor)
    except (ValueError, TypeError):
        return default


def float_seguro(valor, default=0.0):
    try:
        return float(valor)
    except (ValueError, TypeError):
        return default


def bool_seguro(valor, default=False):
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, str):
        return valor.strip().lower() in ["true", "1", "si", "sí", "yes"]
    if isinstance(valor, (int, float)):
        return bool(valor)
    return default


# =========================================================
# FUNCIÓN DE IA (DeepSeek + OpenRouter de respaldo)
# =========================================================
def mejorar_redaccion_ia(texto, tipo_texto="general"):
    """Mejora la redacción usando DeepSeek con respaldo de OpenRouter."""
    if not texto.strip():
        return "Sin información adicional registrada."

    base = "Redacta de forma profesional y clara, sin títulos. No agregar ni quitar información. Mantén esencia y estructura original. Corrige ortografía."

    instrucciones = {
        "reporte": 'Pasado, 3ra persona, un párrafo fluido.',
        "incidente": 'Breve, directo, formal, solo hechos.',
        "nota": 'Tono institucional formal.',
        "resumen": 'Un párrafo fluido, pasado, 3ra persona. Une los hechos narrativamente.',
        "clima": 'Clima técnico.',
        "general": base
    }

    prompt = f"""{base} {instrucciones.get(tipo_texto, "")}

Devuelve SOLO el texto mejorado.

TEXTO:
{texto}

MEJORADO:"""

    if tipo_texto in ["reporte", "resumen", "nota"]:
        max_tok = 1000
    else:
        max_tok = 500

    proveedores = [
        ("https://api.deepseek.com", DEEPSEEK_API_KEY, "deepseek-chat"),
        ("https://openrouter.ai/api/v1", OPENROUTER_API_KEY, "inclusionai/ling-3.0-flash"),
    ]

    for url, key, model in proveedores:
        try:
            client = OpenAI(base_url=url, api_key=key)
            response = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": prompt}],
                max_tokens=max_tok,
                temperature=0.3
            )
            return response.choices[0].message.content.strip()
        except:
            continue

    st.warning("⚠️ Todos los proveedores fallaron. Intenta más tarde.")
    texto_limpio = texto.strip().capitalize()
    if not texto_limpio.endswith('.'):
        texto_limpio += '.'
    return texto_limpio


# =========================================================
# FUNCIÓN DE CLIMA (Windy API)
# =========================================================
def consultar_clima_windy(lat, lon):
    """Consulta el clima actual desde Windy API."""
    url = "https://api.windy.com/api/point-forecast/v2"
    payload = {
        "lat": lat,
        "lon": lon,
        "model": "gfs",
        "parameters": ["wind", "temp", "rh", "pressure", "precip"],
        "levels": ["surface"],
        "key": WINDY_API_KEY,
    }

    try:
        resp = requests.post(url, json=payload, timeout=10)
        if resp.status_code != 200:
            return None, f"Error {resp.status_code}: {resp.text}"

        data = resp.json()

        def safe_get(key, default=0):
            val = data.get(key)
            if val and isinstance(val, list) and len(val) > 0 and val[0] is not None:
                return val[0]
            return default

        temp_k = safe_get("temp-surface", 298.15)
        temp_c = round(temp_k - 273.15)
        humedad = safe_get("rh-surface", 0)
        presion_pa = safe_get("pressure-surface", 101325)
        presion_hpa = round(presion_pa / 100, 4)
        precip_mm = safe_get("precip-surface", 0)

        wind_surface = safe_get("wind-surface", None)
        if wind_surface is not None and wind_surface > 0:
            viento_ms = wind_surface
        else:
            u = safe_get("wind_u-surface", 0)
            v = safe_get("wind_v-surface", 0)
            viento_ms = (u**2 + v**2)**0.5

        viento_kmh = round(float(viento_ms) * 3.6)

        return {
            "viento": f"{viento_kmh:02d} Km/h",
            "temp": f"{temp_c}°C",
            "humedad": f"{round(humedad)}%",
            "presion": f"{presion_hpa:,.4f} hPa".replace(",", "."),
            "precip": f"{int(precip_mm * 10)}%" if precip_mm > 0 else "0%"
        }, None

    except Exception as e:
        return None, str(e)


# =========================================================
# CONSTANTES
# =========================================================
ARCHIVO_MEMORIA = "memoria.json"
ARCHIVO_UBICACIONES = "ubicaciones.json"


# =========================================================
# PERSISTENCIA
# =========================================================
def guardar_memoria(clave, valor):
    memoria = {}
    if os.path.exists(ARCHIVO_MEMORIA):
        try:
            with open(ARCHIVO_MEMORIA, "r", encoding="utf-8") as f:
                memoria = json.load(f)
        except:
            pass
    memoria[clave] = valor
    with open(ARCHIVO_MEMORIA, "w", encoding="utf-8") as f:
        json.dump(memoria, f, ensure_ascii=False, indent=4)


def cargar_memoria(clave, default=""):
    if os.path.exists(ARCHIVO_MEMORIA):
        try:
            with open(ARCHIVO_MEMORIA, "r", encoding="utf-8") as f:
                memoria = json.load(f)
            return memoria.get(clave, default)
        except:
            pass
    return default


# =========================================================
# UBICACIONES GENÉRICAS
# =========================================================
def normalizar_ubicaciones(ubicaciones):
    for region, data_region in ubicaciones.items():
        if not isinstance(data_region, dict):
            ubicaciones[region] = {}
            continue
        ciudades = data_region.get("ciudades", {})
        if not isinstance(ciudades, dict):
            data_region["ciudades"] = {}
            continue
        for ciudad, data_ciudad in ciudades.items():
            if not isinstance(data_ciudad, dict):
                ciudades[ciudad] = {"zonas": [], "barrios": {}, "sub_barrios": {}}
                continue
            if "zonas" not in data_ciudad or not isinstance(data_ciudad["zonas"], list):
                data_ciudad["zonas"] = []
            if "barrios" not in data_ciudad or not isinstance(data_ciudad["barrios"], dict):
                data_ciudad["barrios"] = {}
            if "sub_barrios" not in data_ciudad or not isinstance(data_ciudad["sub_barrios"], dict):
                data_ciudad["sub_barrios"] = {}
            for zona in data_ciudad["zonas"]:
                if zona not in data_ciudad["barrios"]:
                    data_ciudad["barrios"][zona] = []
    return ubicaciones


def cargar_ubicaciones():
    default = {
        "Región Norte": {
            "ciudades": {
                "Ciudad A": {
                    "zonas": ["Zona 1", "Zona 2"],
                    "barrios": {"Zona 1": [], "Zona 2": []},
                    "sub_barrios": {}
                },
                "Ciudad B": {
                    "zonas": ["Zona 3"],
                    "barrios": {"Zona 3": []},
                    "sub_barrios": {}
                }
            }
        },
        "Región Sur": {},
        "Región Este": {},
        "Región Oeste": {}
    }
    if os.path.exists(ARCHIVO_UBICACIONES):
        try:
            with open(ARCHIVO_UBICACIONES, "r", encoding="utf-8") as f:
                data = json.load(f)
            return normalizar_ubicaciones(data)
        except:
            pass
    return normalizar_ubicaciones(default)


def guardar_ubicaciones(ubicaciones):
    try:
        with open(ARCHIVO_UBICACIONES, "w", encoding="utf-8") as f:
            json.dump(ubicaciones, f, ensure_ascii=False, indent=4)
    except:
        pass


# =========================================================
# EXPORTACIÓN A EXCEL BONITO
# =========================================================
COLOR_HEADER = "1F4E3D"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_ROW_ALT = "E8F5E9"
COLOR_TITLE = "28A745"
COLOR_BORDER = "BDBDBD"
COLOR_LABEL = "F1F8E9"


def _estilo_borde():
    lado = Side(style="thin", color=COLOR_BORDER)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _escribir_titulo(ws, fila, texto, num_columnas):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=num_columnas)
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = Font(bold=True, size=16, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=COLOR_TITLE)
    celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[fila].height = 30
    return fila + 1


def _escribir_subtitulo(ws, fila, texto, num_columnas):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=num_columnas)
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = Font(bold=True, size=12, color=COLOR_HEADER)
    celda.fill = PatternFill("solid", fgColor=COLOR_LABEL)
    celda.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[fila].height = 22
    return fila + 1


def _escribir_encabezados(ws, fila, encabezados):
    borde = _estilo_borde()
    for col, texto in enumerate(encabezados, 1):
        celda = ws.cell(row=fila, column=col, value=texto)
        celda.font = Font(bold=True, color=COLOR_HEADER_TEXT, size=11)
        celda.fill = PatternFill("solid", fgColor=COLOR_HEADER)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde
    ws.row_dimensions[fila].height = 24
    return fila + 1


def _escribir_fila_datos(ws, fila, datos, alternar=False):
    borde = _estilo_borde()
    fill = PatternFill("solid", fgColor=COLOR_ROW_ALT) if alternar else None
    for col, valor in enumerate(datos, 1):
        celda = ws.cell(row=fila, column=col, value=valor)
        celda.border = borde
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        if fill:
            celda.fill = fill
    return fila + 1


def _ajustar_anchos(ws, anchos):
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _escribir_pares_clave_valor(ws, fila_inicio, pares, num_columnas=2):
    borde = _estilo_borde()
    fila = fila_inicio
    for clave, valor in pares:
        c_clave = ws.cell(row=fila, column=1, value=clave)
        c_clave.font = Font(bold=True, color=COLOR_HEADER)
        c_clave.fill = PatternFill("solid", fgColor=COLOR_LABEL)
        c_clave.alignment = Alignment(vertical="center", wrap_text=True)
        c_clave.border = borde

        if num_columnas > 1:
            ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=num_columnas)

        c_valor = ws.cell(row=fila, column=2, value=valor)
        c_valor.alignment = Alignment(vertical="center", wrap_text=True)
        c_valor.border = borde
        for col_extra in range(2, num_columnas + 1):
            ws.cell(row=fila, column=col_extra).border = borde

        fila += 1
    return fila


def _escribir_bloque_texto(ws, fila, texto, num_columnas, alto_filas=5):
    """Escribe un bloque de texto multilínea con borde."""
    fila_fin = fila + alto_filas - 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila_fin, end_column=num_columnas)
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.alignment = Alignment(vertical="top", wrap_text=True)
    celda.border = _estilo_borde()
    for r in range(fila, fila_fin + 1):
        for c in range(1, num_columnas + 1):
            ws.cell(row=r, column=c).border = _estilo_borde()
    return fila_fin + 1


def excel_reporte_diario(datos):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Diario"

    _ajustar_anchos(ws, [30, 60])
    fila = 1

    fila = _escribir_titulo(ws, fila, "REPORTE DIARIO", 2)
    fila += 1

    pares = [
        ("N° Reporte", datos.get("num_reporte", "")),
        ("Fecha", datos.get("fecha", "")),
        ("Responsable", datos.get("responsable", "")),
        ("Área", datos.get("area", "")),
        ("Título", datos.get("titulo", "")),
    ]
    fila = _escribir_pares_clave_valor(ws, fila, pares, num_columnas=2)
    fila += 1

    fila = _escribir_subtitulo(ws, fila, "DETALLE", 2)
    fila = _escribir_bloque_texto(ws, fila, datos.get("texto", ""), 2, alto_filas=6)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def excel_reporte_incidente(datos):
    wb = Workbook()
    ws = wb.active
    ws.title = "Incidente"

    _ajustar_anchos(ws, [30, 60])
    fila = 1

    fila = _escribir_titulo(ws, fila, "REPORTE DE INCIDENTE", 2)
    fila += 1

    pares = [
        ("N° Incidente", datos.get("num_inc", "")),
        ("Fecha", datos.get("fecha", "")),
        ("Hora", datos.get("hora", "")),
        ("Tipo", datos.get("tipo", "")),
        ("Reportante", datos.get("reportante", "")),
        ("Ubicación", datos.get("ubicacion", "")),
        ("Coordenadas", datos.get("coordenadas", "")),
    ]
    fila = _escribir_pares_clave_valor(ws, fila, pares, num_columnas=2)
    fila += 1

    fila = _escribir_subtitulo(ws, fila, "DESCRIPCIÓN", 2)
    fila = _escribir_bloque_texto(ws, fila, datos.get("texto", ""), 2, alto_filas=6)
    fila += 1

    fila = _escribir_subtitulo(ws, fila, "CONDICIONES ATMOSFÉRICAS", 2)
    fila = _escribir_encabezados(ws, fila, ["Parámetro", "Valor"])
    clima = [
        ("Viento", datos.get("viento", "")),
        ("Temperatura", datos.get("temp", "")),
        ("Precipitaciones", datos.get("precip", "")),
        ("Humedad", datos.get("humedad", "")),
        ("Presión", datos.get("presion", "")),
    ]
    for i, (k, v) in enumerate(clima):
        fila = _escribir_fila_datos(ws, fila, [k, v], alternar=(i % 2 == 1))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def excel_reporte_mixto(titulo, secciones):
    """secciones = [(nombre, [("k", "v"), ...]) | {"columnas": [...], "filas": [[...], ...]}]"""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:30]

    _ajustar_anchos(ws, [30, 40, 30, 30])
    fila = 1
    fila = _escribir_titulo(ws, fila, titulo.upper(), 4)
    fila += 1

    for nombre, contenido in secciones:
        fila = _escribir_subtitulo(ws, fila, nombre, 4)
        if isinstance(contenido, list):
            fila = _escribir_pares_clave_valor(ws, fila, contenido, num_columnas=4)
        elif isinstance(contenido, dict):
            cols = contenido.get("columnas", [])
            filas = contenido.get("filas", [])
            fila = _escribir_encabezados(ws, fila, cols + [""] * (4 - len(cols)))
            for i, f in enumerate(filas):
                fila = _escribir_fila_datos(ws, fila, f + [""] * (4 - len(f)), alternar=(i % 2 == 1))
        fila += 1

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# =========================================================
# CSS
# =========================================================
st.markdown("""
<style>
    div.stButton > button {
        background: linear-gradient(135deg, #1e7e34, #28a745) !important;
        color: #ffffff !important;
        font-size: 18px !important;
        font-weight: 700 !important;
        padding: 12px 24px !important;
        border-radius: 10px !important;
        border: 2px solid #28a745 !important;
        width: 100% !important;
    }
    div.stButton > button:hover {
        background: linear-gradient(135deg, #155724, #1e7e34) !important;
        transform: scale(1.01) !important;
    }
</style>
""", unsafe_allow_html=True)


# =========================================================
# SIDEBAR
# =========================================================
st.sidebar.title("Sistema de Gestión de Reportes")
st.sidebar.markdown("---")

opcion = st.sidebar.radio(
    "Seleccione un módulo:",
    ["REPORTE DIARIO", "REPORTE DE INCIDENTES", "REPORTES VARIOS"]
)

st.sidebar.markdown("---")
st.sidebar.info("Plantilla genérica con IA + Windy + Excel")


# =========================================================
# MÓDULO 1: REPORTE DIARIO
# =========================================================
if opcion == "REPORTE DIARIO":
    st.header("📋 Reporte Diario")

    col1, col2 = st.columns(2)
    with col1:
        titulo = st.text_input("Título del Reporte", cargar_memoria("diario_titulo", "Reporte General"))
        responsable = st.text_input("Responsable", cargar_memoria("diario_responsable", ""))
        fecha = st.date_input("Fecha", datetime.now(), key="f_diario")
    with col2:
        num_reporte = st.text_input("Número de Reporte", cargar_memoria("diario_num", "001"))
        area = st.text_input("Área / Departamento", cargar_memoria("diario_area", "General"))

    st.subheader("📝 Descripción")
    if "diario_texto" not in st.session_state:
        st.session_state["diario_texto"] = cargar_memoria("diario_texto", "")

    texto = st.text_area(
        "Detalle:",
        placeholder="Describa aquí el reporte...",
        height=150,
        key="diario_texto"
    )

    col_btn1, col_btn2 = st.columns([3, 1])
    with col_btn2:
        if st.button("✨ IA", key="btn_ia_diario"):
            if texto.strip():
                with st.spinner("🤖 Mejorando..."):
                    st.session_state["diario_mejorado"] = mejorar_redaccion_ia(texto, "reporte")
                    st.rerun()
            else:
                st.warning("Escribe algo primero")

    if "diario_mejorado" in st.session_state:
        st.text_area(
            "Versión mejorada:",
            value=st.session_state["diario_mejorado"],
            height=150,
            disabled=True,
            key="diario_mejorado_display"
        )
        c1, c2 = st.columns(2)
        with c1:
            if st.button("✅ Usar mejorado", key="btn_usar_diario"):
                guardar_memoria("diario_texto", st.session_state["diario_mejorado"])
                del st.session_state["diario_texto"]
                del st.session_state["diario_mejorado"]
                st.rerun()
        with c2:
            if st.button("❌ Mantener original", key="btn_mantener_diario"):
                del st.session_state["diario_mejorado"]
                st.rerun()

    if "diario_generado" not in st.session_state:
        st.session_state.diario_generado = ""

    st.markdown("<br>", unsafe_allow_html=True)

    if st.button("📄 GENERAR REPORTE", use_container_width=True):
        if not texto.strip():
            st.warning("⚠️ Escribe el detalle primero.")
        else:
            guardar_memoria("diario_titulo", titulo)
            guardar_memoria("diario_responsable", responsable)
            guardar_memoria("diario_num", num_reporte)
            guardar_memoria("diario_area", area)
            guardar_memoria("diario_texto", texto)

            dias = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
            dia_str = dias[fecha.weekday()]

            st.session_state.diario_generado = f"""*{titulo.upper()}*

*N° REPORTE:* {num_reporte}
*FECHA:* {dia_str} {fecha.strftime('%d/%m/%Y')}
*RESPONSABLE:* {responsable}
*ÁREA:* {area}

*DETALLE:*
{texto}

---"""
            st.success("✅ Reporte generado")

    if st.session_state.diario_generado:
        st.subheader("📋 Reporte Formateado")
        st.code(st.session_state.diario_generado, language=None)

        col_dl1, col_dl2 = st.columns(2)
        with col_dl1:
            excel_buffer = excel_reporte_diario({
                "num_reporte": num_reporte,
                "fecha": fecha.strftime("%d/%m/%Y"),
                "responsable": responsable,
                "area": area,
                "titulo": titulo,
                "texto": texto
            })
            st.download_button(
                label="📥 Descargar Excel",
                data=excel_buffer,
                file_name=f"reporte_diario_{num_reporte or 'sin_num'}_{fecha.strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="dl_diario"
            )
        with col_dl2:
            st.download_button(
                label="📥 Descargar TXT",
                data=st.session_state.diario_generado,
                file_name=f"reporte_diario_{num_reporte or 'sin_num'}.txt",
                mime="text/plain",
                use_container_width=True,
                key="dl_diario_txt"
            )


# =========================================================
# MÓDULO 2: REPORTE DE INCIDENTES
# =========================================================
elif opcion == "REPORTE DE INCIDENTES":
    st.header("⚠️ Reporte de Incidentes")

    if "incidente_texto" not in st.session_state:
        st.session_state["incidente_texto"] = cargar_memoria("incidente_texto", "")

    col1, col2 = st.columns(2)
    with col1:
        num_inc = st.text_input("Número de Incidente", cargar_memoria("inc_num", ""))
        fecha_inc = st.date_input("Fecha", datetime.now(), key="f_inc")
        hora_inc = st.time_input("Hora", datetime.now().time(), key="h_inc")
    with col2:
        opciones_tipo = ["Leve", "Moderado", "Grave", "Crítico"]
        tipo_guardado = cargar_memoria("inc_tipo", "Leve")
        if tipo_guardado not in opciones_tipo:
            tipo_guardado = "Leve"
        tipo_inc = st.selectbox("Tipo de Incidente", opciones_tipo, index=opciones_tipo.index(tipo_guardado))
        reportante = st.text_input("Reportante", cargar_memoria("inc_reportante", ""))

    # Ubicación
    st.subheader("📍 Ubicación")
    ubicaciones = cargar_ubicaciones()
    regiones = list(ubicaciones.keys())

    region_guardada = cargar_memoria("inc_region", regiones[0] if regiones else "")
    if region_guardada not in regiones:
        region_guardada = regiones[0] if regiones else ""

    region = st.selectbox("Región", regiones, index=regiones.index(region_guardada) if regiones else 0)

    if region and ubicaciones[region].get("ciudades"):
        ciudades = list(ubicaciones[region]["ciudades"].keys())
        ciudad_guardada = cargar_memoria("inc_ciudad", ciudades[0] if ciudades else "")
        if ciudad_guardada not in ciudades:
            ciudad_guardada = ciudades[0] if ciudades else ""
        ciudad = st.selectbox("Ciudad", ciudades, index=ciudades.index(ciudad_guardada) if ciudades else 0)

        zonas = ubicaciones[region]["ciudades"][ciudad].get("zonas", [])
        if not zonas:
            zonas = ["(Sin zonas)"]
        zona_guardada = cargar_memoria("inc_zona", zonas[0])
        if zona_guardada not in zonas:
            zona_guardada = zonas[0]
        zona = st.selectbox("Zona", zonas, index=zonas.index(zona_guardada))
    else:
        ciudad = st.text_input("Ciudad", cargar_memoria("inc_ciudad", ""))
        zona = st.text_input("Zona", cargar_memoria("inc_zona", ""))

    ubicacion = f"{zona}, {ciudad}, {region}"

    # Coordenadas + Windy
    st.subheader("📍 Coordenadas y Clima")
    col_c1, col_c2 = st.columns(2)
    with col_c1:
        lat = st.number_input(
            "Latitud",
            value=float_seguro(cargar_memoria("inc_lat", 10.4806), 10.4806),
            format="%.6f",
            key="lat_inc"
        )
    with col_c2:
        lon = st.number_input(
            "Longitud",
            value=float_seguro(cargar_memoria("inc_lon", -66.9036), -66.9036),
            format="%.6f",
            key="lon_inc"
        )

    if st.button("📥 Consultar Clima Automático desde Windy", use_container_width=True):
        with st.spinner("Conectando con Windy..."):
            datos_clima, error = consultar_clima_windy(lat, lon)
            if error:
                st.error(f"No se pudo conectar con Windy: {error}")
            else:
                st.session_state["v_viento"] = datos_clima["viento"]
                st.session_state["v_temp"] = datos_clima["temp"]
                st.session_state["v_hum"] = datos_clima["humedad"]
                st.session_state["v_pres"] = datos_clima["presion"]
                st.session_state["v_precip"] = datos_clima["precip"]
                st.success("✅ Datos atmosféricos sincronizados desde Windy")
                st.rerun()

    col_at1, col_at2, col_at3, col_at4, col_at5 = st.columns(5)
    with col_at1:
        viento_val = st.text_input("Viento (Km/h)", key="v_viento")
    with col_at2:
        temp_val = st.text_input("Temperatura", key="v_temp")
    with col_at3:
        precip_val = st.text_input("Precipitaciones (%)", "0%", key="v_precip")
    with col_at4:
        humedad_val = st.text_input("Humedad", key="v_hum")
    with col_at5:
        presion_val = st.text_input("Presión", key="v_pres")

    st.subheader("📝 Descripción")
    texto_inc = st.text_area(
        "Descripción del incidente:",
        placeholder="Describa el incidente...",
        height=150,
        key="incidente_texto"
    )

    col_btn1, col_btn2 = st.columns([3, 1])
    with col_btn2:
        if st.button("✨ IA", key="btn_ia_inc"):
            if texto_inc.strip():
                with st.spinner("🤖 Mejorando..."):
                    st.session_state["inc_mejorado"] = mejorar_redaccion_ia(texto_inc, "incidente")
                    st.rerun()
            else:
                st.warning("Escribe algo primero")

    if "inc_mejorado" in st.session_state:
        st.text_area(
            "Versión mejorada:",
            value=st.session_state["inc_mejorado"],
            height=150,
            disabled=True,
            key="inc_mejorado_display"
        )
        c1, c2 = st.columns(2)
        with c1:
            if st.button("✅ Usar mejorado", key="btn_usar_inc"):
                guardar_memoria("incidente_texto", st.session_state["inc_mejorado"])
                del st.session_state["incidente_texto"]
                del st.session_state["inc_mejorado"]
                st.rerun()
        with c2:
            if st.button("❌ Mantener original", key="btn_mantener_inc"):
                del st.session_state["inc_mejorado"]
                st.rerun()

    if "inc_generado" not in st.session_state:
        st.session_state.inc_generado = ""

    st.markdown("<br>", unsafe_allow_html=True)

    if st.button("⚠️ GENERAR REPORTE DE INCIDENTE", use_container_width=True):
        if not texto_inc.strip():
            st.warning("⚠️ Escribe la descripción primero.")
        else:
            guardar_memoria("inc_num", num_inc)
            guardar_memoria("inc_tipo", tipo_inc)
            guardar_memoria("inc_reportante", reportante)
            guardar_memoria("inc_region", region)
            guardar_memoria("inc_ciudad", ciudad)
            guardar_memoria("inc_zona", zona)
            guardar_memoria("incidente_texto", texto_inc)
            guardar_memoria("inc_lat", lat)
            guardar_memoria("inc_lon", lon)

            st.session_state.inc_generado = f"""*REPORTE DE INCIDENTE*

*N°:* {num_inc}
*FECHA:* {fecha_inc.strftime('%d/%m/%Y')}
*HORA:* {hora_inc.strftime('%H:%M')}
*TIPO:* {tipo_inc}
*REPORTANTE:* {reportante}
*UBICACIÓN:* {ubicacion}

*DESCRIPCIÓN:*
{texto_inc}

*COORDENADAS:*
{lat}, {lon}

*CONDICIONES ATMOSFÉRICAS:*
✅ Viento: {viento_val}
✅ Temperatura: {temp_val}
✅ Precipitaciones: {precip_val}
✅ Humedad: {humedad_val}
✅ Presión: {presion_val}"""
            st.success("✅ Incidente registrado")

    if st.session_state.inc_generado:
        st.subheader("📋 Reporte Formateado")
        st.code(st.session_state.inc_generado, language=None)

        col_dl1, col_dl2 = st.columns(2)
        with col_dl1:
            excel_buffer = excel_reporte_incidente({
                "num_inc": num_inc,
                "fecha": fecha_inc.strftime("%d/%m/%Y"),
                "hora": hora_inc.strftime("%H:%M"),
                "tipo": tipo_inc,
                "reportante": reportante,
                "ubicacion": ubicacion,
                "coordenadas": f"{lat}, {lon}",
                "texto": texto_inc,
                "viento": viento_val,
                "temp": temp_val,
                "precip": precip_val,
                "humedad": humedad_val,
                "presion": presion_val
            })
            st.download_button(
                label="📥 Descargar Excel",
                data=excel_buffer,
                file_name=f"incidente_{num_inc or 'sin_num'}_{fecha_inc.strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="dl_inc"
            )
        with col_dl2:
            st.download_button(
                label="📥 Descargar TXT",
                data=st.session_state.inc_generado,
                file_name=f"incidente_{num_inc or 'sin_num'}.txt",
                mime="text/plain",
                use_container_width=True,
                key="dl_inc_txt"
            )

    st.markdown("---")

    if st.button("📋 GENERAR REPORTE EJECUTIVO", use_container_width=True, key="btn_ejecutivo_inc"):
        if not texto_inc.strip():
            st.warning("⚠️ Necesitas descripción para generar el ejecutivo.")
        else:
            with st.spinner("🤖 Generando reporte ejecutivo..."):
                texto_combinado = re.sub(r'\d{2}:\d{2}\s*Hrs\s*', '', texto_inc)
                resumen = mejorar_redaccion_ia(texto_combinado, "resumen")

                st.session_state.reporte_ejecutivo_inc = f"""*REPORTE EJECUTIVO*

*FECHA:* {fecha_inc.strftime('%d/%m/%Y')}
*HORA:* {hora_inc.strftime('%H:%M')}
*UBICACIÓN:* {ubicacion}

*EVENTO:* {tipo_inc}

*DESCRIPCIÓN:*
{resumen}

*COORDENADAS:*
{lat},{lon}"""

    if "reporte_ejecutivo_inc" in st.session_state:
        st.subheader("📋 Reporte Ejecutivo Formateado")
        st.code(st.session_state.reporte_ejecutivo_inc, language=None)


# =========================================================
# MÓDULO 3: REPORTES VARIOS
# =========================================================
elif opcion == "REPORTES VARIOS":
    st.header("📄 Reportes Varios")

    tab1, tab2, tab3 = st.tabs(["📌 Nota", "🌤️ Clima", "📊 Resumen"])

    # ----- NOTA -----
    with tab1:
        st.subheader("📝 Nota Informativa")
        fecha_nota = st.date_input("Fecha", datetime.now(), key="f_nota")

        if "nota_texto" not in st.session_state:
            st.session_state["nota_texto"] = cargar_memoria("nota_texto", "")

        texto_nota = st.text_area("Contenido", height=130, key="nota_texto")

        col_ni_btn1, col_ni_btn2 = st.columns([3, 1])
        with col_ni_btn2:
            if st.button("✨ IA", key="btn_ia_nota"):
                if texto_nota.strip():
                    with st.spinner("🤖 Mejorando..."):
                        st.session_state["nota_mejorada"] = mejorar_redaccion_ia(texto_nota, "nota")
                        st.rerun()
                else:
                    st.warning("Escribe algo primero")

        if "nota_mejorada" in st.session_state:
            st.text_area(
                "Nota mejorada:",
                value=st.session_state["nota_mejorada"],
                height=130,
                disabled=True,
                key="nota_mejorada_display"
            )
            c1, c2 = st.columns(2)
            with c1:
                if st.button("✅ Usar mejorado", key="btn_usar_nota"):
                    guardar_memoria("nota_texto", st.session_state["nota_mejorada"])
                    del st.session_state["nota_texto"]
                    del st.session_state["nota_mejorada"]
                    st.rerun()
            with c2:
                if st.button("❌ Mantener original", key="btn_mantener_nota"):
                    del st.session_state["nota_mejorada"]
                    st.rerun()

        autor = st.text_input("Autor", cargar_memoria("nota_autor", ""))

        if st.button("📝 GENERAR NOTA", use_container_width=True, key="btn_nota"):
            if texto_nota.strip():
                dias = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
                dia_str = dias[fecha_nota.weekday()].capitalize()
                guardar_memoria("nota_texto", texto_nota)
                guardar_memoria("nota_autor", autor)

                st.session_state.nota_gen = f"""*NOTA INFORMATIVA*

*FECHA:* {dia_str} {fecha_nota.strftime('%d/%m/%Y')}

{texto_nota}

*RESPONSABLE:* {autor}"""

        if "nota_gen" in st.session_state:
            st.code(st.session_state.nota_gen, language=None)

            excel_buffer = excel_reporte_mixto("NOTA INFORMATIVA", [
                ("DATOS GENERALES", [
                    ("Fecha", fecha_nota.strftime("%d/%m/%Y")),
                    ("Autor", autor),
                ]),
                ("CONTENIDO", {"columnas": ["Texto"], "filas": [[texto_nota]]})
            ])
            st.download_button(
                label="📥 Descargar Excel",
                data=excel_buffer,
                file_name=f"nota_{fecha_nota.strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="dl_nota"
            )

    # ----- CLIMA -----
    with tab2:
        st.subheader("🌤️ Reporte Meteorológico")

        col_m1, col_m2 = st.columns(2)
        with col_m1:
            region_met = st.text_input("Región", cargar_memoria("met_region", "Región Norte"))
            ciudad_met = st.text_input("Ciudad", cargar_memoria("met_ciudad", "Ciudad A"))
            fecha_met = st.date_input("Fecha", datetime.now(), key="f_met_rep")
        with col_m2:
            hora_met = st.text_input("Hora", cargar_memoria("met_hora", "07:26 Hrs"))
            capacidad = st.number_input(
                "Capacidad Operativa",
                min_value=0,
                value=int_seguro(cargar_memoria("met_capacidad", 20), 20),
                step=1
            )

        if "met_condiciones" not in st.session_state:
            st.session_state["met_condiciones"] = cargar_memoria("met_condiciones", "Cielo despejado, vientos moderados.")

        condiciones_met = st.text_area("Condiciones Atmosféricas", height=80, key="met_condiciones")

        col_met_btn1, col_met_btn2 = st.columns([3, 1])
        with col_met_btn2:
            if st.button("✨ IA", key="btn_ia_cond_met"):
                if condiciones_met.strip():
                    with st.spinner("🤖 Mejorando..."):
                        st.session_state["cond_mejorada_met"] = mejorar_redaccion_ia(condiciones_met, "clima")
                        st.rerun()

        if "cond_mejorada_met" in st.session_state:
            st.text_area(
                "Condiciones mejoradas:",
                value=st.session_state["cond_mejorada_met"],
                height=80,
                disabled=True,
                key="cond_mejorada_display_met"
            )
            c1, c2 = st.columns(2)
            with c1:
                if st.button("✅ Usar mejorado", key="btn_usar_cond_met"):
                    guardar_memoria("met_condiciones", st.session_state["cond_mejorada_met"])
                    del st.session_state["met_condiciones"]
                    del st.session_state["cond_mejorada_met"]
                    st.rerun()
            with c2:
                if st.button("❌ Mantener original", key="btn_mantener_cond_met"):
                    del st.session_state["cond_mejorada_met"]
                    st.rerun()

        if "met_acciones" not in st.session_state:
            st.session_state["met_acciones"] = cargar_memoria("met_acciones", "Personal en alerta preventiva.")

        acciones_met = st.text_area("Acciones Realizadas", height=80, key="met_acciones")

        if st.button("🌤️ GENERAR REPORTE METEOROLÓGICO", use_container_width=True, key="btn_rm"):
            dias = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
            dia_str = dias[fecha_met.weekday()].capitalize()

            guardar_memoria("met_region", region_met)
            guardar_memoria("met_ciudad", ciudad_met)
            guardar_memoria("met_hora", hora_met)
            guardar_memoria("met_capacidad", capacidad)
            guardar_memoria("met_condiciones", condiciones_met)
            guardar_memoria("met_acciones", acciones_met)

            st.session_state.reporte_met = f"""*REPORTE METEOROLÓGICO*

*REGIÓN:* {region_met}
*CIUDAD:* {ciudad_met}
*FECHA:* {dia_str} {fecha_met.strftime('%d/%m/%Y')}
*HORA:* {hora_met}

*CAPACIDAD OPERATIVA:* {capacidad}

*CONDICIONES ATMOSFÉRICAS:* {condiciones_met}

*ACCIONES REALIZADAS:*
{acciones_met}"""

        if "reporte_met" in st.session_state:
            st.code(st.session_state.reporte_met, language=None)

            excel_buffer = excel_reporte_mixto("REPORTE METEOROLÓGICO", [
                ("DATOS GENERALES", [
                    ("Región", region_met),
                    ("Ciudad", ciudad_met),
                    ("Fecha", fecha_met.strftime("%d/%m/%Y")),
                    ("Hora", hora_met),
                    ("Capacidad Operativa", str(capacidad)),
                ]),
                ("CONDICIONES", {"columnas": ["Parámetro", "Valor"], "filas": [
                    ["Condiciones Atmosféricas", condiciones_met],
                    ["Acciones Realizadas", acciones_met],
                ]})
            ])
            st.download_button(
                label="📥 Descargar Excel",
                data=excel_buffer,
                file_name=f"meteo_{fecha_met.strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="dl_met"
            )

    # ----- RESUMEN -----
    with tab3:
        st.subheader("📊 Resumen de Memoria")
        if os.path.exists(ARCHIVO_MEMORIA):
            try:
                with open(ARCHIVO_MEMORIA, "r", encoding="utf-8") as f:
                    datos = json.load(f)

                # Tabla bonita
                if datos:
                    filas = [{"Clave": k, "Valor": str(v)[:80]} for k, v in datos.items()]
                    st.dataframe(filas, use_container_width=True, hide_index=True)

                    # Botón Excel del resumen completo
                    excel_buffer = excel_reporte_mixto("RESUMEN GENERAL", [
                        ("DATOS GUARDADOS", {
                            "columnas": ["Clave", "Valor"],
                            "filas": [[k, str(v)] for k, v in datos.items()]
                        })
                    ])
                    st.download_button(
                        label="📥 Descargar Resumen en Excel",
                        data=excel_buffer,
                        file_name=f"resumen_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        key="dl_resumen"
                    )
                else:
                    st.info("La memoria está vacía.")
            except:
                st.info("No se pudo leer la memoria.")
        else:
            st.info("No hay datos guardados aún.")

# FOOTER
st.markdown("---")
st.caption("Sistema de Gestión de Reportes - Plantilla Genérica con IA + Windy + Excel")
